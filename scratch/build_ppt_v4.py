"""
Build PPT v4 - Gaya Ma Chung (sesuai referensi teman)
Style: Latar krem/putih bersih, judul besar, ikon teal, minimal whitespace,
       logo Ma Chung + Sistem Informasi di pojok kiri atas.
"""

from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree
import os, copy

# ============================================================
# Paths
# ============================================================
BASE = r'c:\xampp\htdocs\inventory_zencare\lain\File PPT'
DEST = os.path.join(BASE, 'Presentasi_Sempro_Marcell.pptx')
IMG = {
    'erd':           os.path.join(BASE, r'ERD Logical Database\Entity Relationship Diagram (ERD) & Logical Database.png'),
    'usecase':       os.path.join(BASE, r'Use Case\TA Use Case.png'),
    'sitemap_p':     os.path.join(BASE, r'Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical - Pelanggan.png'),
    'sitemap_a':     os.path.join(BASE, r'Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical - Admin.png'),
    'sitemap_sa':    os.path.join(BASE, r'Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical - Superadmin.png'),
    'act_terima':    os.path.join(BASE, r'Activity Diagram\Kelola Penerimaan Barang & Validasi Logistik.png'),
    'act_pos':       os.path.join(BASE, r'Activity Diagram\Pesanan Luring (Transaksi Kasir (POS)).png'),
    'act_ecom':      os.path.join(BASE, r'Activity Diagram\Pesanan Daring (E-Commerce).png'),
    'act_pesanan':   os.path.join(BASE, r'Activity Diagram\Kelola Pesanan Daring & Pembatalan.png'),
    'wf_keluar_b':   os.path.join(BASE, r'Workflow\Workflow Barang Keluar (Berjalan).png'),
    'wf_keluar_s':   os.path.join(BASE, r'Workflow\Workflow Barang Keluar (Sistem Usulan).png'),
    'wf_masuk_b':    os.path.join(BASE, r'Workflow\Workflow Barang Masuk (Berjalan).png'),
    'wf_masuk_s':    os.path.join(BASE, r'Workflow\Workflow Barang Masuk & Logistik (Sistem Usulan).png'),
    'wf_pos':        os.path.join(BASE, r'Wireframe\POS.png'),
    'wf_dashboard_a':os.path.join(BASE, r'Wireframe\Dashboard (Admin).png'),
    'wf_dashboard_s':os.path.join(BASE, r'Wireframe\Dashboard (Superadmin).png'),
    'wf_penerimaan': os.path.join(BASE, r'Wireframe\Penerimaan Barang.png'),
    'wf_kartu':      os.path.join(BASE, r'Wireframe\Kartu Stok.png'),
    'wf_transaksi':  os.path.join(BASE, r'Wireframe\Transaksi.png'),
    'wf_ecom':       os.path.join(BASE, r'Wireframe\E-Commerce.png'),
    'wf_checkout':   os.path.join(BASE, r'Wireframe\Checkout Pembayaran.png'),
    'wf_lapjual':    os.path.join(BASE, r'Wireframe\Laporan Penjualan Barang.png'),
    'wf_laplogistik':os.path.join(BASE, r'Wireframe\Laporan Logistik Medis.png'),
    'wf_lappers':    os.path.join(BASE, r'Wireframe\Laporan Persediaan.png'),
    'wf_mutasi':     os.path.join(BASE, r'Wireframe\Mutasi Stok.png'),
}

# ============================================================
# Color Palette (Ma Chung style - referensi teman)
# ============================================================
BG          = RGBColor(0xF4, 0xF2, 0xED)  # Cream/beige background
TEAL        = RGBColor(0x4A, 0xB5, 0xBF)  # Ma Chung teal (icons, accents)
TEAL_DARK   = RGBColor(0x2D, 0x89, 0x96)  # Darker teal (table header, etc.)
TITLE_COLOR = RGBColor(0x1A, 0x1A, 0x2E)  # Near black for titles
TEXT_DARK   = RGBColor(0x2D, 0x3A, 0x4A)  # Dark text
TEXT_GRAY   = RGBColor(0x6B, 0x7A, 0x8D)  # Muted gray text
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
RED_X       = RGBColor(0xE5, 0x35, 0x35)
GREEN_CHK   = RGBColor(0x2E, 0x9E, 0x5E)
TABLE_HDR   = RGBColor(0x4A, 0xB5, 0xBF)  # Same as teal
TABLE_ALT   = RGBColor(0xE8, 0xF5, 0xF6)  # Light teal row

# ============================================================
# Slide size: 13.33" x 7.5" (16:9)
# ============================================================
W, H = 13.33, 7.5
def e(v): return int(v * 914400)

# ============================================================
# Low-level helpers
# ============================================================
def new_prs():
    prs = Presentation()
    prs.slide_width  = e(W)
    prs.slide_height = e(H)
    return prs

def blank(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # Fill background with cream
    bg = s.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = BG
    return s

def tb(slide, l, t, w, h, text, sz=14, bold=False, color=TEXT_DARK,
       align=PP_ALIGN.LEFT, wrap=True, italic=False):
    box = slide.shapes.add_textbox(e(l), e(t), e(w), e(h))
    tf  = box.text_frame
    tf.word_wrap = wrap
    p   = tf.paragraphs[0]
    p.alignment = align
    r   = p.add_run()
    r.text        = text
    r.font.size   = Pt(sz)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box

def mtb(slide, l, t, w, h, lines, base_sz=14, wrap=True):
    """Multi-paragraph textbox: lines = list of (text, sz, bold, color, align)"""
    box = slide.shapes.add_textbox(e(l), e(t), e(w), e(h))
    tf  = box.text_frame
    tf.word_wrap = wrap
    for i, (text, sz, bold, color, align) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        r = p.add_run()
        r.text = text; r.font.size = Pt(sz)
        r.font.bold = bold; r.font.color.rgb = color
    return box

def rect(slide, l, t, w, h, fill=None, line=None, lw=0.75):
    s = slide.shapes.add_shape(1, e(l), e(t), e(w), e(h))
    if fill:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    return s

def img(slide, path, l, t, w, h):
    if os.path.exists(path):
        slide.shapes.add_picture(path, e(l), e(t), e(w), e(h))

def circle(slide, cx, cy, r_in, fill=TEAL):
    """Draw filled circle centered at (cx,cy) with radius r_in."""
    x = cx - r_in; y = cy - r_in
    s = slide.shapes.add_shape(9, e(x), e(y), e(r_in*2), e(r_in*2))  # 9=OVAL
    s.fill.solid(); s.fill.fore_color.rgb = fill
    s.line.fill.background()
    return s

# ============================================================
# Standard decorators (logo area + title + underline)
# ============================================================
def logo_area(slide):
    """Top-left: Ma Chung logo placeholder + Program Studi label."""
    rect(slide, 0.28, 0.18, 1.30, 0.52, fill=WHITE)
    tb(slide, 0.30, 0.20, 1.26, 0.48, 'MA CHUNG\nUniversitas', sz=7, bold=True,
       color=TEAL_DARK, align=PP_ALIGN.CENTER)
    rect(slide, 1.65, 0.18, 1.20, 0.52, fill=WHITE)
    tb(slide, 1.67, 0.20, 1.16, 0.48, 'Program Studi\nSistem Informasi',
       sz=7, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

def logo_area_topright(slide):
    """Top-right: Ma Chung logo placeholder."""
    rect(slide, W - 2.90, 0.12, 1.30, 0.52, fill=WHITE)
    tb(slide, W - 2.88, 0.14, 1.26, 0.48, 'MA CHUNG\nUniversitas', sz=7, bold=True,
       color=TEAL_DARK, align=PP_ALIGN.CENTER)
    rect(slide, W - 1.50, 0.12, 1.25, 0.52, fill=WHITE)
    tb(slide, W - 1.48, 0.14, 1.20, 0.48, 'Program Studi\nSistem Informasi',
       sz=7, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

def page_title_left(slide, title, with_logo_left=True, with_logo_right=False):
    """Ma Chung style: Large title top-left, teal underline, logo area."""
    if with_logo_left:
        logo_area(slide)
    if with_logo_right:
        logo_area_topright(slide)
    # Title
    tb(slide, 0.50, 0.85, W - 1.00, 0.85,
       title, sz=34, bold=True, color=TITLE_COLOR, align=PP_ALIGN.LEFT)
    # Short teal underline (like referensi)
    rect(slide, 0.50, 1.75, 0.65, 0.05, fill=TEAL)

def page_title_center(slide, title, with_logo_left=True):
    """Center-title variant (slide Siklus Hidup style)."""
    if with_logo_left:
        logo_area(slide)
    tb(slide, 0.50, 0.85, W - 1.00, 0.85,
       title, sz=34, bold=True, color=TITLE_COLOR, align=PP_ALIGN.CENTER)
    rect(slide, W/2 - 0.40, 1.75, 0.80, 0.05, fill=TEAL)

# ============================================================
# Icon above text helper (Ma Chung style - teal emoji icon)
# ============================================================
def icon_col(slide, x, y, icon, label, body, icon_sz=28, lbl_sz=14, body_sz=12, col_w=2.80):
    """Single column: icon → bold label → body text (Ma Chung minimal style)."""
    tb(slide, x, y, col_w, 0.55, icon, sz=icon_sz, color=TEAL, align=PP_ALIGN.LEFT)
    tb(slide, x, y + 0.60, col_w, 0.40, label, sz=lbl_sz, bold=True, color=TEXT_DARK)
    tb(slide, x, y + 1.05, col_w, 2.00, body, sz=body_sz, color=TEXT_GRAY)

def icon_col_circle(slide, cx, y, icon, label, body, col_w=2.00):
    """Icon in teal circle (Siklus Hidup style)."""
    r = 0.42
    circle(slide, cx + col_w/2, y + r, r, fill=TEAL)
    tb(slide, cx, y, col_w, r*2, icon, sz=22, color=WHITE, align=PP_ALIGN.CENTER)
    tb(slide, cx, y + r*2 + 0.12, col_w, 0.36, label, sz=12, bold=True,
       color=TEXT_DARK, align=PP_ALIGN.CENTER)
    tb(slide, cx, y + r*2 + 0.52, col_w, 0.80, body, sz=10, color=TEXT_GRAY,
       align=PP_ALIGN.CENTER)

# ============================================================
# SLIDE BUILDERS
# ============================================================

def s01_cover(prs):
    s = blank(prs)
    # Full teal left stripe
    rect(s, 0, 0, 0.30, H, fill=TEAL)
    # Bottom dark band
    rect(s, 0, H - 1.10, W, 1.10, fill=TEAL_DARK)

    # Logo (top left above stripe)
    tb(s, 0.50, 0.25, 2.50, 0.45, 'Universitas Ma Chung', sz=11, bold=True,
       color=TEAL_DARK)
    tb(s, 0.50, 0.72, 2.50, 0.30, 'Program Studi Sistem Informasi', sz=9, color=TEXT_GRAY)

    # Title block - big and clean
    tb(s, 0.50, 1.45, W - 0.80, 1.90,
       'RANCANG BANGUN\nSISTEM INFORMASI INVENTORY OMNICHANNEL\nALAT KESEHATAN DAN OBAT',
       sz=30, bold=True, color=TITLE_COLOR)
    rect(s, 0.50, 3.42, 0.80, 0.06, fill=TEAL)
    tb(s, 0.50, 3.58, W - 0.80, 0.45,
       'Studi Kasus: ZenCare Medical',
       sz=17, italic=True, color=TEXT_GRAY)

    # Author
    rect(s, 0.50, 4.25, W - 0.80, 0.04, fill=RGBColor(0xCC, 0xCC, 0xCC))
    tb(s, 0.50, 4.40, W - 0.80, 0.40,
       'Marcell Chandra Kenchana  —  NIM: 322310015', sz=15, bold=True, color=TEXT_DARK)
    tb(s, 0.50, 4.85, W - 0.80, 0.35,
       'Program Studi Sistem Informasi  •  Fakultas Teknologi dan Desain  •  Universitas Ma Chung', sz=11, color=TEXT_GRAY)

    # Pembimbing area (bottom band)
    tb(s, 0.50, H - 1.00, W/2 - 0.30, 0.88,
       'Dosen Pembimbing:\nDr. Soetam Rizky Wicaksono, S.Kom., MM., MCP., MCTS., MOSM.',
       sz=11, bold=False, color=WHITE)
    tb(s, W/2 + 0.20, H - 1.00, W/2 - 0.50, 0.88,
       'Penguji:\n[Ketua Penguji]  |  [Penguji 1]  |  [Penguji 2]',
       sz=11, color=WHITE)
    tb(s, W - 2.20, 0.25, 2.00, 0.40, 'Malang, 2026', sz=11, italic=True,
       color=TEXT_GRAY, align=PP_ALIGN.RIGHT)


def s02_toc(prs):
    s = blank(prs)
    logo_area(s)
    page_title_center(s, 'Daftar Isi', with_logo_left=False)
    logo_area(s)

    items = [
        ('📌', 'Latar Belakang &\nIdentifikasi Masalah', 'Masalah overselling,\nkonversi satuan,\nlogistik medis'),
        ('📚', 'Penelitian Terdahulu\n& Research Gap', 'Perbandingan &\nkebaruan sistem\nZenCare'),
        ('🔄', 'Workflow & Use Case\nSistem', 'Alur bisnis berjalan\nvs. sistem usulan\nomnichannel'),
        ('🗄️', 'Perancangan Sistem\n& Database', 'Sitemap, Activity Diagram,\nERD, Logical Database,\nWireframe'),
        ('✅', 'Kesimpulan\n& Luaran', 'Simpulan penelitian\n& luaran yang\ndiharapkan'),
    ]
    cw = (W - 1.00) / 5
    for i, (icon, label, body) in enumerate(items):
        x = 0.50 + i * (cw + 0.00)
        icon_col_circle(s, x, 2.20, icon, label, body, col_w=cw - 0.10)


def s03_latbel(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Latar Belakang', with_logo_left=False, with_logo_right=True)

    cards = [
        ('🏪', 'Kasir Luring & E-Commerce\nTidak Terhubung',
         'Aplikasi POS pihak ketiga bersifat stand-alone — sisa stok tidak tersinkronisasi otomatis ke e-commerce. '
         'Risiko overselling sangat tinggi: barang habis di toko masih bisa dipesan secara daring.'),
        ('🔢', 'Konversi Multi-Satuan\nManual (Multi-UOM)',
         'E-commerce berbasis satuan grosir (Box), kasir melayani eceran (Strip/Pcs). '
         'Admin harus menghitung konversi kemasan secara manual setiap transaksi daring — rentan human error.'),
        ('💊', 'Ketiadaan Kendali\nLogistik Medis',
         'Tanggal kedaluwarsa obat tidak didata per Batch (FEFO). Alat Kesehatan tidak dicatat '
         'per Serial Number (SN) untuk validasi garansi. Tidak ada stiker Batch → pengambilan tidak tervalidasi sistem.'),
    ]
    col_w = (W - 1.00) / 3
    for i, (icon, label, body) in enumerate(cards):
        x = 0.50 + i * col_w
        icon_col(s, x, 2.05, icon, label, body, icon_sz=32, lbl_sz=14, body_sz=12, col_w=col_w - 0.25)


def s04_batasan(prs):
    s = blank(prs)
    logo_area(s)
    page_title_center(s, 'Batasan Masalah', with_logo_left=False)
    logo_area(s)

    items = [
        ('Fokus Operasional',
         'Sinkronisasi inventaris antara POS dan E-Commerce mandiri. '
         'Tidak mencakup akuntansi keuangan (HPP/Laba-Rugi), sistem poin pelanggan, maupun auto-replenishment stok.'),
        ('Logika Pengeluaran Logistik',
         'Menerapkan First Expired First Out (FEFO) berbasis Nomor Batch untuk obat. '
         'Pencatatan Serial Number (SN) per unit untuk Alat Kesehatan. Validasi data wajib pada penerimaan barang.'),
        ('Infrastruktur Pihak Ketiga',
         'E-Commerce dibangun mandiri (bukan marketplace eksternal). Payment Gateway Midtrans & RajaOngkir '
         'diuji pada lingkungan Sandbox. Beroperasi pada satu lokasi: ZenCare Medical Muharto.'),
        ('Cakupan Fitur',
         'Tidak ada pre-order, retur daring otomatis, atau refund otomatis. '
         'Penyesuaian stok akibat retur/garansi dikelola manual melalui modul Mutasi Stok.'),
    ]
    col_w = (W - 1.00) / 4
    for i, (label, body) in enumerate(items):
        x = 0.50 + i * col_w
        tb(s, x, 2.05, col_w - 0.20, 0.35, label, sz=13, bold=True, color=TEXT_DARK)
        rect(s, x, 2.43, col_w - 0.30, 0.04, fill=TEAL)
        tb(s, x, 2.55, col_w - 0.20, 4.50, body, sz=12, color=TEXT_GRAY)


def s05_gap(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Perbandingan Penelitian Terdahulu', with_logo_left=False, with_logo_right=True)

    # Table
    cols  = ['Aspek', 'Rachman (2024)\n& Nuralisa (2024)', 'Mattegunta (2025)', '★ Sistem Usulan\n(ZenCare Medical)']
    col_w = [1.55, 3.20, 3.20, 4.58]
    col_x = [0.35]
    for w in col_w[:-1]:
        col_x.append(col_x[-1] + w)
    rows = [
        ('Penulis\n& Tahun', 'Rachman, Hartanto,\nHadiwijaya (2024)\nNuralisa, Nirsal (2024)', 'Mattegunta\n(2025)', 'Marcell Chandra K.\n(2026)'),
        ('Platform', 'Web (React.js /\nPHP Laravel)\nSingle-channel', 'POS + Cloud\nEkosistem ritel umum', 'PHP + MySQL\nOmnichannel mandiri'),
        ('Sinkronisasi', '✗  Tidak ada\n(stand-alone)', '✔  Real-time\n(ritel umum)', '✔  Real-time\nBasis data tunggal'),
        ('Konversi\nMulti-UOM', '✗  Tidak ada', '✗  Tidak dibahas', '✔  Dual-UOM otomatis\nBox ↔ Strip'),
        ('Logistik\nMedis', '✗  Tidak ada', '✗  Tidak dibahas', '✔  FEFO Batch (Obat)\n+ SN (Alat Kesehatan)'),
    ]

    th = 0.62
    rh = 0.95
    ty = 1.90
    for ci, (hdr, cw, cx) in enumerate(zip(cols, col_w, col_x)):
        c = TEAL_DARK if ci < 3 else TEAL
        rect(s, cx, ty, cw, th, fill=c)
        tb(s, cx + 0.06, ty + 0.06, cw - 0.12, th - 0.10,
           hdr, sz=11, bold=True, color=WHITE, align=PP_ALIGN.CENTER)

    for ri, row in enumerate(rows):
        ry = ty + th + ri * rh
        for ci, (cell, cw, cx) in enumerate(zip(row, col_w, col_x)):
            is_last = ri == len(rows) - 1
            bg = TABLE_ALT if ri % 2 == 0 else WHITE
            if ci == 3 and is_last:
                bg = RGBColor(0xD4, 0xED, 0xDA)
            elif is_last and ci > 0:
                bg = RGBColor(0xFF, 0xEC, 0xEC)
            if ci == 0:
                bg = RGBColor(0xE0, 0xF4, 0xF5)
            rect(s, cx, ry, cw, rh, fill=bg, line=RGBColor(0xCC, 0xD5, 0xD8), lw=0.5)
            clr = RED_X if '✗' in cell else (GREEN_CHK if '✔' in cell else TEXT_DARK)
            bd  = ci == 0 or (ci == 3 and is_last)
            tb(s, cx + 0.06, ry + 0.08, cw - 0.12, rh - 0.14,
               cell, sz=11, bold=bd, color=clr)


def s06_novelty(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Pembaruan Sistem (Novelty)', with_logo_left=False, with_logo_right=True)

    # Left header
    tb(s, 0.50, 2.00, 5.80, 0.30, 'PENELITIAN TERDAHULU', sz=9, bold=True, color=TEXT_GRAY)
    rect(s, 0.50, 2.32, 5.80, 0.03, fill=RGBColor(0xCC, 0xCC, 0xCC))
    probs = [
        'Sistem bersifat single-channel (luring saja) tanpa sinkronisasi omnichannel.',
        'Tidak ada fitur konversi multi-satuan (Multi-UOM) otomatis lintas saluran.',
        'Tidak ada pencatatan tanggal kedaluwarsa berbasis Batch (FEFO) untuk obat.',
        'Tidak ada pelacakan Serial Number (SN) untuk validasi garansi alat kesehatan.',
    ]
    for i, p in enumerate(probs):
        y = 2.50 + i * 1.00
        tb(s, 0.50, y, 0.35, 0.40, '✕', sz=16, bold=True, color=RED_X)
        tb(s, 0.92, y + 0.03, 5.20, 0.80, p, sz=13, color=TEXT_DARK)

    # Divider
    rect(s, 6.55, 1.90, 0.04, 5.30, fill=RGBColor(0xCC, 0xCC, 0xCC))

    # Right header
    tb(s, 6.80, 2.00, 5.80, 0.30, 'SISTEM USULAN (ZENCARE MEDICAL)', sz=9, bold=True, color=TEXT_GRAY)
    rect(s, 6.80, 2.32, 5.80, 0.03, fill=TEAL)
    sols = [
        'Basis data tunggal Omnichannel: POS & E-Commerce terhubung, stok terkunci seketika (Reserved).',
        'Algoritma Dual-UOM mengonversi stok lintas satuan (Box ↔ Strip) secara otomatis tanpa hitungan manual.',
        'Sistem Gatekeeper memaksa input Nomor Batch + Tanggal Kadaluwarsa sebelum stok masuk — logika FEFO.',
        'Setiap unit Alat Kesehatan terdaftar per SN di tabel unit_serial → riwayat garansi dapat ditelusuri.',
    ]
    for i, p in enumerate(sols):
        y = 2.50 + i * 1.00
        tb(s, 6.80, y, 0.35, 0.40, '✓', sz=16, bold=True, color=GREEN_CHK)
        tb(s, 7.20, y + 0.03, 5.75, 0.80, p, sz=13, color=TEXT_DARK)


def s07_usecase(prs):
    s = blank(prs)
    logo_area(s)
    page_title_left(s, 'Use Case Diagram', with_logo_left=False)
    logo_area(s)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Pemodelan hak akses dan interaksi 3 aktor: Pelanggan, Admin (Staff), dan Superadmin (Pemilik).',
       sz=11, italic=True, color=TEXT_GRAY)
    img(s, IMG['usecase'], 0.50, 2.20, W - 1.00, H - 2.50)


def s08_workflow_before(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Workflow Bisnis Berjalan', with_logo_left=False, with_logo_right=True)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Kondisi saat ini: sistem kasir berdiri sendiri, konversi stok dan sinkronisasi e-commerce dilakukan secara manual.',
       sz=11, italic=True, color=TEXT_GRAY)
    half = (W - 1.20) / 2
    tb(s, 0.50, 2.22, half, 0.30, '📤  Barang Keluar (Berjalan)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['wf_keluar_b'], 0.50, 2.55, half, H - 2.80)
    rx = 0.50 + half + 0.20
    tb(s, rx, 2.22, half, 0.30, '📥  Barang Masuk (Berjalan)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['wf_masuk_b'], rx, 2.55, half, H - 2.80)


def s09_workflow_after(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Workflow Sistem Usulan — Omnichannel', with_logo_left=False, with_logo_right=True)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Sistem usulan: satu basis data terpusat — pemotongan stok otomatis dengan FEFO, tidak ada konversi manual.',
       sz=11, italic=True, color=TEXT_GRAY)
    half = (W - 1.20) / 2
    tb(s, 0.50, 2.22, half, 0.30, '📤  Barang Keluar (Sistem Usulan)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['wf_keluar_s'], 0.50, 2.55, half, H - 2.80)
    rx = 0.50 + half + 0.20
    tb(s, rx, 2.22, half, 0.30, '📥  Barang Masuk & Logistik (Sistem Usulan)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['wf_masuk_s'], rx, 2.55, half, H - 2.80)


def s10_sitemap(prs):
    s = blank(prs)
    logo_area(s)
    page_title_left(s, 'Sitemap — Arsitektur Navigasi Sistem', with_logo_left=False)
    logo_area(s)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Hak akses dipisahkan menjadi 3 lingkungan: Pelanggan (Front-End E-Commerce), Admin (Operasional), Superadmin (Manajerial).',
       sz=11, italic=True, color=TEXT_GRAY)
    third = (W - 1.20) / 3
    for i, (key, lbl, clr) in enumerate([
        ('sitemap_p',  '👤  Pelanggan', TEAL),
        ('sitemap_a',  '🧑‍💼  Admin', TEAL_DARK),
        ('sitemap_sa', '👑  Superadmin', RGBColor(0x2C, 0x6E, 0x8A)),
    ]):
        x = 0.50 + i * (third + 0.10)
        tb(s, x, 2.22, third, 0.30, lbl, sz=12, bold=True, color=clr)
        img(s, IMG[key], x, 2.55, third, H - 2.80)


def s11_act_terima(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Activity Diagram — Penerimaan Barang', with_logo_left=False, with_logo_right=True)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Sistem bertindak sebagai Gatekeeper: identitas logistik medis (Batch/SN) wajib diisi sebelum stok diproses.',
       sz=11, italic=True, color=TEXT_GRAY)
    img(s, IMG['act_terima'], 0.50, 2.20, W - 1.00, H - 2.50)


def s12_act_pos_ecom(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Activity Diagram — POS & E-Commerce', with_logo_left=False, with_logo_right=True)
    half = (W - 1.20) / 2
    tb(s, 0.50, 1.85, half, 0.28, '🏪  Transaksi Kasir (POS)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['act_pos'], 0.50, 2.18, half, H - 2.45)
    rx = 0.50 + half + 0.20
    tb(s, rx, 1.85, half, 0.28, '🌐  Pesanan Daring (E-Commerce)', sz=12, bold=True, color=TEAL_DARK)
    img(s, IMG['act_ecom'], rx, 2.18, half, H - 2.45)


def s13_erd(prs):
    s = blank(prs)
    logo_area(s)
    page_title_left(s, 'Entity Relationship Diagram (ERD) & Logical Database', with_logo_left=False)
    logo_area(s)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       '13 tabel relasional  |  4 klaster: Aktor/RBAC · Multi-UOM · Logistik Medis · Transaksi & Audit Trail',
       sz=11, italic=True, color=TEXT_GRAY)
    img(s, IMG['erd'], 0.50, 2.20, W - 1.00, H - 2.50)


def s14_wireframe(prs):
    s = blank(prs)
    logo_area(s)
    page_title_left(s, 'Wireframe — Rancangan Antarmuka Utama', with_logo_left=False)
    logo_area(s)
    tb(s, 0.50, 1.85, W - 1.00, 0.30,
       'Rancangan struktural antarmuka: fokus pada fungsionalitas dan tata letak, bukan estetika visual akhir.',
       sz=11, italic=True, color=TEXT_GRAY)

    screens = [
        ('wf_pos',        '🏪 POS (Kasir)'),
        ('wf_dashboard_a','📊 Dashboard Admin'),
        ('wf_penerimaan', '📥 Penerimaan Barang'),
        ('wf_ecom',       '🌐 E-Commerce (Beranda)'),
        ('wf_checkout',   '💳 Checkout Pembayaran'),
        ('wf_laplogistik','📋 Laporan Logistik Medis'),
    ]
    sw = (W - 1.20) / 3
    sh = (H - 2.50) / 2
    for i, (key, label) in enumerate(screens):
        col = i % 3
        row = i // 3
        x = 0.50 + col * (sw + 0.10)
        y = 2.22 + row * (sh + 0.06)
        tb(s, x, y, sw, 0.25, label, sz=10, bold=True, color=TEAL_DARK)
        img(s, IMG[key], x, y + 0.28, sw, sh - 0.30)


def s15_kesimpulan(prs):
    s = blank(prs)
    logo_area_topright(s)
    page_title_left(s, 'Kesimpulan & Luaran Penelitian', with_logo_left=False, with_logo_right=True)

    # Simpulan block (clean, no heavy box)
    tb(s, 0.50, 2.00, W - 1.00, 0.25, 'SIMPULAN', sz=9, bold=True, color=TEXT_GRAY)
    rect(s, 0.50, 2.28, W - 1.00, 0.03, fill=RGBColor(0xCC, 0xCC, 0xCC))
    tb(s, 0.50, 2.38, W - 1.00, 1.70,
       'Dirancang sistem informasi inventori omnichannel berbasis web untuk ZenCare Medical yang '
       'menyinkronkan transaksi POS dan E-Commerce dalam satu basis data terpusat menggunakan metode RAD. '
       'Sistem mengotomatisasi konversi Multi-UOM, menerapkan Gatekeeper FEFO untuk obat berbasis Batch, '
       'dan mencatat Serial Number untuk garansi Alat Kesehatan — menciptakan tata kelola logistik medis '
       'yang akurat, aman, dan transparan.',
       sz=13.5, color=TEXT_DARK)

    # Luaran section
    tb(s, 0.50, 4.20, W - 1.00, 0.25, 'LUARAN PENELITIAN', sz=9, bold=True, color=TEXT_GRAY)
    rect(s, 0.50, 4.48, W - 1.00, 0.03, fill=RGBColor(0xCC, 0xCC, 0xCC))

    luaran = [
        ('🌐', 'Sistem Omnichannel\nBerbasis Web',
         'POS + E-Commerce terintegrasi dalam satu basis data tunggal.'),
        ('🔄', 'Otomatisasi\nDual-UOM',
         'Konversi Box ↔ Strip berjalan otomatis, tanpa hitungan manual kasir.'),
        ('💊', 'Gatekeeper FEFO\n& Serial Number',
         'Validasi logistik medis wajib: Batch, kadaluwarsa, dan SN per unit.'),
        ('📋', 'Dokumentasi\n& Laporan UAT',
         'Laporan pengujian sinkronisasi stok & evaluasi penerimaan pengguna.'),
    ]
    cw = (W - 1.00) / 4
    for i, (icon, lbl, body) in enumerate(luaran):
        x = 0.50 + i * cw
        icon_col(s, x, 4.65, icon, lbl, body, icon_sz=24, lbl_sz=12, body_sz=11, col_w=cw - 0.15)


def s16_penutup(prs):
    s = blank(prs)
    rect(s, 0, 0, 0.30, H, fill=TEAL)
    rect(s, 0, H - 0.85, W, 0.85, fill=TEAL_DARK)

    logo_area(s)

    tb(s, 0.50, 2.00, W - 0.70, 1.50,
       'Terima Kasih', sz=56, bold=True, color=TITLE_COLOR, align=PP_ALIGN.CENTER)
    rect(s, W/2 - 0.80, 3.58, 1.60, 0.06, fill=TEAL)
    tb(s, 0.50, 3.75, W - 0.70, 0.60,
       'Sesi Tanya Jawab  (Q&A)', sz=22, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

    tb(s, 0.50, 4.65, W - 0.70, 0.35,
       'Marcell Chandra Kenchana  •  NIM: 322310015', sz=14, bold=True, color=TEXT_DARK, align=PP_ALIGN.CENTER)
    tb(s, 0.50, 5.05, W - 0.70, 0.32,
       'Program Studi Sistem Informasi  •  Universitas Ma Chung', sz=11, color=TEXT_GRAY, align=PP_ALIGN.CENTER)

    tb(s, 0.50, H - 0.78, W - 0.70, 0.65,
       'Dosen Pembimbing: Dr. Soetam Rizky Wicaksono, S.Kom., MM., MCP., MCTS., MOSM.   •   Malang, 2026',
       sz=10, color=WHITE, align=PP_ALIGN.CENTER)


# ============================================================
# MAIN
# ============================================================
def build():
    prs = new_prs()
    builders = [
        (s01_cover,        'Cover'),
        (s02_toc,          'Daftar Isi'),
        (s03_latbel,       'Latar Belakang'),
        (s04_batasan,      'Batasan Masalah'),
        (s05_gap,          'Research Gap / Tabel'),
        (s06_novelty,      'Novelty'),
        (s07_usecase,      'Use Case Diagram'),
        (s08_workflow_before, 'Workflow Berjalan'),
        (s09_workflow_after,  'Workflow Usulan'),
        (s10_sitemap,      'Sitemap'),
        (s11_act_terima,   'Activity Diagram Penerimaan'),
        (s12_act_pos_ecom, 'Activity Diagram POS & E-Com'),
        (s13_erd,          'ERD'),
        (s14_wireframe,    'Wireframe'),
        (s15_kesimpulan,   'Kesimpulan'),
        (s16_penutup,      'Penutup'),
    ]
    for fn, name in builders:
        fn(prs)
        print('  [' + str(len(prs.slides)) + '/16] ' + name)

    prs.save(DEST)
    import os as _os
    sz = _os.path.getsize(DEST)
    print()
    print('SUCCESS: ' + DEST)
    print('Size: ' + str(round(sz/1024)) + ' KB  |  Slides: ' + str(len(prs.slides)))

if __name__ == '__main__':
    build()
