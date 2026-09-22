"""
Script: build_sempro_canva.py
Desain: Canva-style Ma Chung (cream background, clean typography, no heavy sidebar)
Referensi: template teman (latar belakang cream, title bold top-left, logo top-right, accent teal)
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os

# ─── COLOR PALETTE ────────────────────────────────────────────────────────────
TEAL       = RGBColor(0x4A, 0xB5, 0xBF)   # Ma Chung teal
TEAL_DARK  = RGBColor(0x2D, 0x89, 0x96)   # darker teal
TEAL_LIGHT = RGBColor(0xD4, 0xF0, 0xF3)   # very light teal for accents
CREAM      = RGBColor(0xF5, 0xF2, 0xEB)   # cream background (like the photo)
DARK       = RGBColor(0x2D, 0x3A, 0x40)   # charcoal for body text
TITLE_CLR  = RGBColor(0x1A, 0x2A, 0x30)   # near-black for titles
GRAY       = RGBColor(0x7A, 0x8A, 0x90)   # muted gray for subtitles
DIVIDER    = RGBColor(0xC8, 0xD8, 0xDC)   # light divider
RED_MUTED  = RGBColor(0xC0, 0x39, 0x2B)
GREEN_MUTED = RGBColor(0x27, 0x7A, 0x55)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)

# ─── PATHS ────────────────────────────────────────────────────────────────────
BASE = r'C:\xampp\htdocs\inventory_zencare\lain\File PPT'
IMG = {
    'sitemap_admin'      : rf'{BASE}\Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical (Admin).png',
    'sitemap_superadmin' : rf'{BASE}\Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical (Superadmin).png',
    'sitemap_pelanggan'  : rf'{BASE}\Sitemap\Sitemap Sistem Informasi Inventori Omnichannel ZenCare Medical (Pelanggan).png',
    'erd'                : rf'{BASE}\ERD Logical Database\Entity Relationship Diagram (ERD) & Logical Database.png',
    'usecase'            : rf'{BASE}\Use Case\TA Use Case.png',
    'workflow_keluar_berjalan' : rf'{BASE}\Workflow\Workflow Barang Keluar (Berjalan).png',
    'workflow_keluar_usulan'   : rf'{BASE}\Workflow\Workflow Barang Keluar (Sistem Usulan).png',
    'workflow_masuk_berjalan'  : rf'{BASE}\Workflow\Workflow Barang Masuk (Berjalan).png',
    'workflow_masuk_usulan'    : rf'{BASE}\Workflow\Workflow Barang Masuk & Logistik (Sistem Usulan).png',
    'wf_penerimaan'      : rf'{BASE}\Wireframe\Penerimaan Barang.png',
    'wf_dashboard_sa'    : rf'{BASE}\Wireframe\Dashboard (Superadmin).png',
    'wf_pos'             : rf'{BASE}\Wireframe\POS.png',
    'wf_checkout'        : rf'{BASE}\Wireframe\Checkout Pembayaran.png',
    'activity_logistik'  : rf'{BASE}\Activity Diagram\Kelola Penerimaan Barang & Validasi Logistik.png',
    'activity_keluar'    : rf'{BASE}\Activity Diagram\Pesanan Luring (Transaksi Kasir (POS)).png',
}
OUTPUT = rf'{BASE}\Sempro_ZenCare_Canva.pptx'

# ─── PRESENTATION SETUP ───────────────────────────────────────────────────────
W = Inches(13.33)
H = Inches(7.50)
prs = Presentation()
prs.slide_width  = W
prs.slide_height = H
blank_layout = prs.slide_layouts[6]


# ─── HELPERS ──────────────────────────────────────────────────────────────────
def new_slide():
    return prs.slides.add_slide(blank_layout)


def fill_bg(slide, color=CREAM):
    """Fill entire slide background"""
    bg = slide.shapes.add_shape(1, 0, 0, W, H)
    bg.fill.solid()
    bg.fill.fore_color.rgb = color
    bg.line.fill.background()
    return bg


def add_rect(slide, l, t, w, h, color, line=False):
    shape = slide.shapes.add_shape(1, l, t, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    if not line:
        shape.line.fill.background()
    return shape


def tb(slide, l, t, w, h, text, size=12, bold=False, color=DARK,
       align=PP_ALIGN.LEFT, italic=False, wrap=True):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = wrap
    para = tf.paragraphs[0]
    para.alignment = align
    run = para.add_run()
    run.text = text
    run.font.size   = Pt(size)
    run.font.bold   = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return box


def add_img(slide, path, l, t, w, h):
    if os.path.exists(path):
        slide.shapes.add_picture(path, l, t, w, h)
    else:
        ph = add_rect(slide, l, t, w, h, TEAL_LIGHT)
        tb(slide, l+Inches(0.05), t+h//2, w, Inches(0.4),
           f'[{os.path.basename(path)}]', size=7, color=GRAY, align=PP_ALIGN.CENTER)


def slide_header(slide, title, subtitle=None):
    """Canva-style header: title top-left, teal accent line, logo top-right"""
    # Title
    tb(slide, Inches(0.65), Inches(0.30), Inches(9.0), Inches(0.75),
       title, size=28, bold=True, color=TITLE_CLR)
    # Teal underline accent (short, like in the photo)
    add_rect(slide, Inches(0.65), Inches(1.08), Inches(0.80), Inches(0.055), TEAL)
    if subtitle:
        tb(slide, Inches(0.65), Inches(1.18), Inches(9.0), Inches(0.40),
           subtitle, size=12, color=GRAY, italic=True)

    # Top-right logos placeholder area (text labels since we don't have logo images)
    add_rect(slide, Inches(11.20), Inches(0.15), Inches(1.95), Inches(0.65), TEAL_LIGHT)
    tb(slide, Inches(11.22), Inches(0.18), Inches(0.90), Inches(0.55),
       'MA CHUNG', size=7, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)
    add_rect(slide, Inches(11.25), Inches(0.40), Inches(0.85), Inches(0.02), TEAL_DARK)
    tb(slide, Inches(12.18), Inches(0.18), Inches(0.90), Inches(0.55),
       'Sistem\nInformasi', size=6.5, color=TEAL_DARK, align=PP_ALIGN.CENTER)


def content_top():
    """Y position for content start"""
    return Inches(1.40)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 1: COVER
# ═══════════════════════════════════════════════════════════════════════════════
s1 = new_slide()
fill_bg(s1, CREAM)

# Decorative top-left teal block (like Ma Chung logo area)
add_rect(s1, Inches(0.50), Inches(0.40), Inches(0.04), Inches(0.90), TEAL)

# Logo area top-left
add_rect(s1, Inches(0.65), Inches(0.35), Inches(2.50), Inches(0.55), TEAL_LIGHT)
tb(s1, Inches(0.70), Inches(0.38), Inches(1.20), Inches(0.50), 'MA CHUNG', size=8, bold=True, color=TEAL_DARK)
tb(s1, Inches(2.00), Inches(0.40), Inches(1.00), Inches(0.45), 'Sistem\nInformasi', size=6, color=TEAL_DARK)

# "SEMINAR PROPOSAL TUGAS AKHIR" label
tb(s1, Inches(0.65), Inches(1.50), Inches(11.0), Inches(0.40),
   'SEMINAR PROPOSAL TUGAS AKHIR', size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER)

# Main title — centered, large
title_box = s1.shapes.add_textbox(Inches(0.65), Inches(1.95), Inches(12.0), Inches(1.80))
tf = title_box.text_frame
tf.word_wrap = True
for line, txt in enumerate([
    'Rancang Bangun Sistem Informasi',
    'Inventory Omnichannel',
    'Alat Kesehatan dan Obat'
]):
    p = tf.paragraphs[0] if line == 0 else tf.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(34)
    r.font.bold = True
    r.font.color.rgb = TITLE_CLR

# Studi kasus
tb(s1, Inches(0.65), Inches(3.85), Inches(12.0), Inches(0.40),
   'Studi Kasus: ZenCare Medical', size=14, italic=True, color=GRAY, align=PP_ALIGN.CENTER)

# Teal divider line (center)
add_rect(s1, Inches(5.50), Inches(4.38), Inches(2.35), Inches(0.055), TEAL)

# Name & NIM
tb(s1, Inches(0.65), Inches(4.50), Inches(12.0), Inches(0.45),
   'Marcell Chandra Kenchana', size=15, bold=True, color=DARK, align=PP_ALIGN.CENTER)
tb(s1, Inches(0.65), Inches(4.92), Inches(12.0), Inches(0.35),
   '322310015  —  Program Studi Sistem Informasi', size=11, color=GRAY, align=PP_ALIGN.CENTER)

# Penguji info
tb(s1, Inches(0.65), Inches(5.45), Inches(12.0), Inches(0.32),
   'Ketua Penguji  :  [Nama Ketua Penguji]', size=11, color=DARK, align=PP_ALIGN.CENTER)
tb(s1, Inches(0.65), Inches(5.78), Inches(12.0), Inches(0.32),
   'Dosen Penguji I  :  [Nama Dosen Penguji I]', size=11, color=DARK, align=PP_ALIGN.CENTER)
tb(s1, Inches(0.65), Inches(6.11), Inches(12.0), Inches(0.32),
   'Dosen Penguji II  :  Dr. Soetam Rizky Wicaksono, S.Kom., MM., MCP., MCTS., MOSM.', size=11, color=DARK, align=PP_ALIGN.CENTER)

# Bottom thin line
add_rect(s1, Inches(0.50), Inches(6.85), Inches(12.35), Inches(0.06), TEAL_DARK)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 2: DAFTAR ISI
# ═══════════════════════════════════════════════════════════════════════════════
s2 = new_slide()
fill_bg(s2, CREAM)
slide_header(s2, 'Daftar Isi')

ct = content_top()
items = [
    ('01', 'Latar Belakang &\nIdentifikasi Masalah'),
    ('02', 'Perbandingan\nPenelitian Terdahulu'),
    ('03', 'Workflow &\nAntarmuka Sistem'),
    ('04', 'Arsitektur Teknis\n& Database'),
    ('05', 'Dokumentasi Lapangan\n& Kesimpulan'),
]
bw = Inches(2.30)
gap = Inches(0.18)
total = bw * 5 + gap * 4
sx = (W - total) / 2
for i, (num, label) in enumerate(items):
    bx = sx + i * (bw + gap)
    by = ct + Inches(0.60)
    # clean card: just teal top border + white background
    add_rect(s2, bx, by, bw, Inches(2.40), WHITE)
    add_rect(s2, bx, by, bw, Inches(0.06), TEAL)
    tb(s2, bx, by + Inches(0.18), bw, Inches(0.70),
       num, size=30, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    add_rect(s2, bx + Inches(0.4), by + Inches(0.95), bw - Inches(0.8), Inches(0.04), DIVIDER)
    tb(s2, bx + Inches(0.05), by + Inches(1.10), bw - Inches(0.1), Inches(1.15),
       label, size=12, color=DARK, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 3: LATAR BELAKANG (Canva icon-style)
# ═══════════════════════════════════════════════════════════════════════════════
s3 = new_slide()
fill_bg(s3, CREAM)
slide_header(s3, 'Latar Belakang & Identifikasi Masalah')

ct = content_top() + Inches(0.35)
cols = [
    ('📦', 'Kasir Luring & Daring Terpisah',
     'Penggunaan aplikasi POS fisik stand-alone memicu risiko pesanan ganda (overselling) karena ketersediaan fisik tidak tersinkronisasi dengan etalase e-commerce.'),
    ('🔢', 'Konversi Multi-Satuan Manual',
     'Ketiadaan fitur konversi satuan memaksa admin menghitung pecahan kemasan grosir (Box) ke eceran (Strip) secara manual sebelum memotong sisa persediaan.'),
    ('🏥', 'Ketiadaan Kendali Logistik Medis',
     'Sistem pencatatan saat ini belum mendata tanggal kedaluwarsa berbasis Batch (FEFO) untuk obat, maupun Serial Number (SN) jaminan garansi alat kesehatan.'),
]
cw = Inches(3.80)
gap = Inches(0.28)
total = cw * 3 + gap * 2
sx = (W - total) / 2
for i, (icon, title, body) in enumerate(cols):
    cx = sx + i * (cw + gap)
    # Icon (large, teal)
    tb(s3, cx, ct, cw, Inches(0.72),
       icon, size=34, color=TEAL, align=PP_ALIGN.LEFT)
    # Title bold
    tb(s3, cx, ct + Inches(0.78), cw, Inches(0.50),
       title, size=14, bold=True, color=TITLE_CLR)
    # Thin divider under title
    add_rect(s3, cx, ct + Inches(1.33), Inches(0.55), Inches(0.04), TEAL)
    # Body text
    tb(s3, cx, ct + Inches(1.50), cw - Inches(0.1), Inches(3.20),
       body, size=12, color=GRAY)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 4: BATASAN MASALAH (clean 3-column, Canva style)
# ═══════════════════════════════════════════════════════════════════════════════
s4 = new_slide()
fill_bg(s4, CREAM)
slide_header(s4, 'Batasan Masalah')

ct = content_top() + Inches(0.45)
batasan = [
    ('🎯', 'Fokus Operasional',
     'Sistem murni menyinkronkan inventaris antara terminal POS dan E-Commerce mandiri. Tidak mencakup modul akuntansi keuangan atau laba rugi.'),
    ('📋', 'Logika Pengeluaran Logistik',
     'Pengendalian persediaan menerapkan algoritma First Expired First Out (FEFO) berbasis Batch untuk obat, serta pendataan Serial Number untuk Alat Kesehatan.'),
    ('🔗', 'Infrastruktur Pihak Ketiga',
     'Integrasi Midtrans API (pembayaran) dan RajaOngkir (ongkir) dijalankan secara eksklusif pada lingkungan pengujian (Sandbox Mode).'),
]
cw = Inches(3.80)
gap = Inches(0.28)
sx = (W - (cw * 3 + gap * 2)) / 2
for i, (icon, title, body) in enumerate(batasan):
    cx = sx + i * (cw + gap)
    # Vertical left teal accent bar
    add_rect(s4, cx, ct, Inches(0.05), Inches(3.80), TEAL)
    tb(s4, cx + Inches(0.20), ct, Inches(0.55), Inches(0.65),
       icon, size=28, color=TEAL)
    tb(s4, cx + Inches(0.20), ct + Inches(0.65), cw - Inches(0.25), Inches(0.50),
       title, size=14, bold=True, color=TITLE_CLR)
    tb(s4, cx + Inches(0.20), ct + Inches(1.22), cw - Inches(0.30), Inches(2.50),
       body, size=12, color=GRAY)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 5: RESEARCH GAP — Clean table (Canva style)
# ═══════════════════════════════════════════════════════════════════════════════
s5 = new_slide()
fill_bg(s5, CREAM)
slide_header(s5, 'Perbandingan Penelitian Terdahulu')

ct = content_top() + Inches(0.10)
headers = ['Aspek', 'Rachman dkk.\n(2024)', 'Nuralisa dkk.\n(2024)', 'Mattegunta\n(2025)', 'Sistem Usulan\n(ZenCare Medical)']
rows_data = [
    ['Lingkup Saluran',        'Single-channel\n(Luring)', 'Single-channel\n(Luring)', 'Omnichannel\nRitel Umum', '✔ Omnichannel\n(POS + E-Commerce)'],
    ['Konversi Multi-UOM',     '✗ Tidak Ada', '✗ Tidak Ada', '✗ Tidak Ada', '✔ Dual-UOM Dinamis'],
    ['Pelacakan FEFO/Batch',   '✗ Tidak Ada', '✗ Tidak Ada', '✗ Tidak Ada', '✔ FEFO per Batch'],
    ['Serial Number Alkes',    '✗ Tidak Ada', '✗ Tidak Ada', '✗ Tidak Ada', '✔ SN Alat Kesehatan'],
    ['Segmentasi Domain',      '✗ Umum',      '✗ Umum',      '✗ Ritel Umum', '✔ Medis (Obat & Alkes)'],
]
col_w = [Inches(2.40), Inches(1.98), Inches(1.98), Inches(1.98), Inches(2.50)]
row_h = Inches(0.67)
tx = Inches(0.60)
ty = ct

# Header row
hx = tx
for j, (hdr, cw_) in enumerate(zip(headers, col_w)):
    c = TEAL_DARK if j == 0 else TEAL
    add_rect(s5, hx, ty, cw_, row_h, c)
    tb(s5, hx+Inches(0.06), ty+Inches(0.08), cw_-Inches(0.12), row_h-Inches(0.10),
       hdr, size=10, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    hx += cw_

for ri, row in enumerate(rows_data):
    ry = ty + (ri+1) * row_h
    rx = tx
    for j, (cell, cw_) in enumerate(zip(row, col_w)):
        bg = RGBColor(0xEB, 0xF7, 0xF8) if ri % 2 == 0 else WHITE
        if j == 4:
            bg = RGBColor(0xE4, 0xF7, 0xED)
        if j == 0:
            bg = RGBColor(0xD6, 0xEC, 0xEE)
        add_rect(s5, rx, ry, cw_, row_h, bg)
        tc = DARK
        if '✔' in cell:
            tc = GREEN_MUTED
        elif '✗' in cell:
            tc = RED_MUTED
        bold_ = j == 0
        tb(s5, rx+Inches(0.06), ry+Inches(0.06), cw_-Inches(0.12), row_h-Inches(0.08),
           cell, size=10, color=tc, align=PP_ALIGN.CENTER, bold=bold_)
        rx += cw_


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 6: NOVELTY — Canva clean split (X vs V, no heavy boxes)
# ═══════════════════════════════════════════════════════════════════════════════
s6 = new_slide()
fill_bg(s6, CREAM)
slide_header(s6, 'Pembaruan Sistem (Novelty)')

ct = content_top() + Inches(0.30)
half = Inches(5.90)
mid  = Inches(0.45) + half + Inches(0.45)

# Left: Before (red accent)
tb(s6, Inches(0.65), ct, half, Inches(0.38),
   'PENELITIAN TERDAHULU', size=10, bold=True, color=RED_MUTED)
before = [
    'Risiko overselling akibat aplikasi kasir yang terisolasi dari stok E-Commerce.',
    'Beban kerja tinggi akibat konversi satuan (Box→Strip) dilakukan manual oleh admin.',
    'Penumpukan obat kadaluarsa di rak karena tidak ada sistem peringatan FEFO.',
    'Nomor Seri alkes tidak terdokumentasi, menyulitkan klaim garansi.',
    'Tidak ada pemisahan wewenang (RBAC) antara staf gudang dan manajemen.',
]
by_ = ct + Inches(0.52)
for item in before:
    tb(s6, Inches(0.65), by_, Inches(0.35), Inches(0.45),
       '✗', size=16, bold=True, color=RED_MUTED)
    tb(s6, Inches(1.05), by_ + Inches(0.03), half - Inches(0.50), Inches(0.45),
       item, size=11.5, color=DARK)
    by_ += Inches(0.88)

# Vertical divider
add_rect(s6, Inches(6.65), ct, Inches(0.04), Inches(5.20), DIVIDER)

# Right: After (green accent)
tb(s6, Inches(6.90), ct, half, Inches(0.38),
   'SISTEM USULAN (ZENCARE MEDICAL)', size=10, bold=True, color=GREEN_MUTED)
after = [
    'Sinkronisasi basis data tunggal Omnichannel mengunci stok seketika di semua saluran.',
    'Algoritma konversi Dual-UOM dinamis (rasio_konversi) berjalan otomatis.',
    'Sistem Gatekeeper: validasi FEFO wajib di setiap transaksi barang masuk & keluar.',
    'Setiap unit Alat Kesehatan teregistrasi Serial Number unik di basis data.',
    'RBAC tegas: Admin mencatat fisik, Superadmin mengontrol Master Produk & Harga.',
]
ay_ = ct + Inches(0.52)
for item in after:
    tb(s6, Inches(6.90), ay_, Inches(0.35), Inches(0.45),
       '✔', size=16, bold=True, color=GREEN_MUTED)
    tb(s6, Inches(7.30), ay_ + Inches(0.03), half - Inches(0.50), Inches(0.45),
       item, size=11.5, color=DARK)
    ay_ += Inches(0.88)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 7: SITEMAP
# ═══════════════════════════════════════════════════════════════════════════════
s7 = new_slide()
fill_bg(s7, CREAM)
slide_header(s7, 'Arsitektur Navigasi (Sitemap)')

ct = content_top() + Inches(0.10)
sw = Inches(4.10)
gap = Inches(0.18)
sx = (W - (sw * 3 + gap * 2)) / 2

for i, (key, label) in enumerate([
    ('sitemap_pelanggan', 'Pelanggan (E-Commerce)'),
    ('sitemap_admin', 'Admin (Staf Gudang & Kasir)'),
    ('sitemap_superadmin', 'Superadmin (Manajemen)'),
]):
    cx = sx + i * (sw + gap)
    # Label bar
    add_rect(s7, cx, ct, sw, Inches(0.42), TEAL if i == 2 else TEAL_LIGHT)
    tb(s7, cx + Inches(0.05), ct + Inches(0.06), sw - Inches(0.1), Inches(0.35),
       label, size=11, bold=True, color=TEAL_DARK if i < 2 else WHITE, align=PP_ALIGN.CENTER)
    # Image
    add_img(s7, IMG[key], cx, ct + Inches(0.45), sw, Inches(4.40))

# Caption
tb(s7, Inches(0.65), ct + Inches(5.05), Inches(12.0), Inches(0.42),
   'Superadmin mengontrol Master Produk & Harga serta memantau riwayat pengadaan dari Dashboard. Admin mengelola POS & Penerimaan Barang. Pelanggan berinteraksi via E-Commerce.',
   size=10, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 8: WORKFLOW BARANG KELUAR
# ═══════════════════════════════════════════════════════════════════════════════
s8 = new_slide()
fill_bg(s8, CREAM)
slide_header(s8, 'Workflow Barang Keluar (Omnichannel)')

ct = content_top() + Inches(0.10)
wh = Inches(5.90)
gap = Inches(0.40)

# Label + Workflow images top
tb(s8, Inches(0.65), ct, wh, Inches(0.30), 'Sistem Berjalan', size=10, bold=False, color=GRAY)
add_img(s8, IMG['workflow_keluar_berjalan'], Inches(0.65), ct + Inches(0.30), wh, Inches(2.30))
tb(s8, Inches(0.65) + wh + gap, ct, wh, Inches(0.30), 'Sistem Usulan', size=10, bold=True, color=TEAL_DARK)
add_img(s8, IMG['workflow_keluar_usulan'], Inches(0.65) + wh + gap, ct + Inches(0.30), wh, Inches(2.30))

# Thin divider
add_rect(s8, Inches(0.65), ct + Inches(2.75), Inches(12.0), Inches(0.04), DIVIDER)

# Wireframes bottom
bbot = ct + Inches(2.90)
wfw = Inches(5.90)
tb(s8, Inches(0.65), bbot, wfw, Inches(0.28), 'Wireframe: Terminal POS (Transaksi Luring)', size=10, color=GRAY)
add_img(s8, IMG['wf_pos'], Inches(0.65), bbot + Inches(0.28), wfw, Inches(2.00))
tb(s8, Inches(0.65) + wfw + gap, bbot, wfw, Inches(0.28), 'Wireframe: Checkout E-Commerce (Transaksi Daring)', size=10, color=GRAY)
add_img(s8, IMG['wf_checkout'], Inches(0.65) + wfw + gap, bbot + Inches(0.28), wfw, Inches(2.00))

# Caption bar
tb(s8, Inches(0.65), Inches(6.32), Inches(12.0), Inches(0.38),
   '🔄  Sistem mengonversi satuan otomatis dan mengunci stok sementara (Reserved) untuk mencegah overselling di semua saluran secara bersamaan.',
   size=10, color=GRAY, italic=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 9: WORKFLOW PENERIMAAN BARANG
# ═══════════════════════════════════════════════════════════════════════════════
s9 = new_slide()
fill_bg(s9, CREAM)
slide_header(s9, 'Workflow Penerimaan Barang & Logistik')

ct = content_top() + Inches(0.10)
wh = Inches(5.90)
gap = Inches(0.40)

tb(s9, Inches(0.65), ct, wh, Inches(0.30), 'Sistem Berjalan', size=10, color=GRAY)
add_img(s9, IMG['workflow_masuk_berjalan'], Inches(0.65), ct + Inches(0.30), wh, Inches(2.30))
tb(s9, Inches(0.65) + wh + gap, ct, wh, Inches(0.30), 'Sistem Usulan', size=10, bold=True, color=TEAL_DARK)
add_img(s9, IMG['workflow_masuk_usulan'], Inches(0.65) + wh + gap, ct + Inches(0.30), wh, Inches(2.30))

add_rect(s9, Inches(0.65), ct + Inches(2.75), Inches(12.0), Inches(0.04), DIVIDER)

bbot = ct + Inches(2.90)
wfw = Inches(5.90)
tb(s9, Inches(0.65), bbot, wfw, Inches(0.28), 'Wireframe: Form Penerimaan Barang (Admin)', size=10, color=GRAY)
add_img(s9, IMG['wf_penerimaan'], Inches(0.65), bbot + Inches(0.28), wfw, Inches(2.00))
tb(s9, Inches(0.65) + wfw + gap, bbot, wfw, Inches(0.28), 'Wireframe: Dashboard Superadmin', size=10, color=GRAY)
add_img(s9, IMG['wf_dashboard_sa'], Inches(0.65) + wfw + gap, bbot + Inches(0.28), wfw, Inches(2.00))

tb(s9, Inches(0.65), Inches(6.32), Inches(12.0), Inches(0.38),
   '🔐  Admin wajib menginput Identitas Logistik (Batch/EXP/SN) & Harga Beli. Superadmin memantau histori pengadaan & harga modal via Dashboard secara real-time.',
   size=10, color=GRAY, italic=True)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 10: USE CASE DIAGRAM
# ═══════════════════════════════════════════════════════════════════════════════
s10 = new_slide()
fill_bg(s10, CREAM)
slide_header(s10, 'Analisis Kebutuhan Sistem (Use Case Diagram)')

ct = content_top() + Inches(0.05)
add_img(s10, IMG['usecase'], Inches(1.50), ct, Inches(10.35), Inches(5.10))
tb(s10, Inches(0.65), ct + Inches(5.20), Inches(12.0), Inches(0.38),
   'Tiga aktor: Pelanggan (E-Commerce)  ·  Admin Gudang/Kasir (POS & Logistik)  ·  Superadmin (Master Produk, Harga & Laporan Manajerial).',
   size=10, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 11: ACTIVITY DIAGRAM
# ═══════════════════════════════════════════════════════════════════════════════
s11 = new_slide()
fill_bg(s11, CREAM)
slide_header(s11, 'Activity Diagram: Penerimaan Barang & Validasi Logistik')

ct = content_top() + Inches(0.05)
img_w = Inches(7.80)
img_l = (W - img_w) / 2
add_img(s11, IMG['activity_logistik'], img_l, ct, img_w, Inches(5.05))

# Explanation card right side
note_x = img_l + img_w + Inches(0.20)
note_w = W - note_x - Inches(0.45)
if note_w > Inches(2.0):
    add_rect(s11, note_x, ct, note_w, Inches(5.05), TEAL_LIGHT)
    add_rect(s11, note_x, ct, Inches(0.05), Inches(5.05), TEAL)
    tb(s11, note_x + Inches(0.15), ct + Inches(0.20), note_w - Inches(0.20), Inches(0.45),
       'Poin Kunci', size=11, bold=True, color=TEAL_DARK)
    notes = [
        'Sistem memvalidasi kelengkapan data logistik sebelum menyimpan.',
        'Obat → wajib No. Batch + Tanggal Kedaluwarsa (EXP).',
        'Alkes → wajib Serial Number (SN).',
        'Jika tidak lengkap, proses dikembalikan ke form input.',
        'Harga Beli (Modal) wajib diisi oleh Admin.',
    ]
    ny = ct + Inches(0.75)
    for note in notes:
        tb(s11, note_x + Inches(0.15), ny, note_w - Inches(0.25), Inches(0.78),
           f'• {note}', size=10, color=DARK)
        ny += Inches(0.83)

tb(s11, Inches(0.65), ct + Inches(5.18), Inches(12.0), Inches(0.38),
   'Alur percabangan logika if/else memastikan data logistik medis tidak pernah kosong sebelum stok diperbarui.',
   size=10, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 12: ERD LOGICAL DATABASE
# ═══════════════════════════════════════════════════════════════════════════════
s12 = new_slide()
fill_bg(s12, CREAM)
slide_header(s12, 'Logical Database: Entity Relationship Diagram (ERD)')

ct = content_top() + Inches(0.05)
add_img(s12, IMG['erd'], Inches(0.65), ct, Inches(12.0), Inches(5.10))

add_rect(s12, Inches(0.65), Inches(6.30), Inches(12.0), Inches(0.04), TEAL_LIGHT)
tb(s12, Inches(0.65), Inches(6.40), Inches(12.0), Inches(0.38),
   'Arsitektur Relasional 4 Klaster:  Klaster Aktor (RBAC)  ·  Konversi Multi-UOM  ·  Logistik Medis (Stok & Batch)  ·  Transaksi & Riwayat Mutasi (Kartu Stok)',
   size=10, color=GRAY, italic=True, align=PP_ALIGN.CENTER)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 13: STRUKTUR TABEL ESENSIAL
# ═══════════════════════════════════════════════════════════════════════════════
s13 = new_slide()
fill_bg(s13, CREAM)
slide_header(s13, 'Struktur Tabel Database Esensial')

ct = content_top() + Inches(0.20)
half = Inches(5.85)
gap  = Inches(0.30)
sx   = Inches(0.65)

tables = [
    {
        'title': 'produk_variasi — Konversi Multi-UOM',
        'desc': 'Mengatur konversi satuan dinamis dan harga jual statis per varian produk.',
        'fields': [
            ('rasio_konversi', 'INT', 'Pengunci perkalian matematis grosir → eceran'),
            ('satuan_kecil / satuan_besar', 'VARCHAR', 'Contoh: Pcs / Box'),
            ('harga_jual_kecil / besar', 'DECIMAL', 'Harga jual statis dikontrol Superadmin'),
            ('tampil_di_online', 'TINYINT', 'Toggle visibilitas E-Commerce per varian'),
        ]
    },
    {
        'title': 'stok_batch — Logika FEFO Obat',
        'desc': 'Menyimpan nomor batch dan tanggal kadaluarsa; sistem memotong stok berdasarkan FEFO.',
        'fields': [
            ('no_batch', 'VARCHAR', 'Nomor batch dari faktur supplier'),
            ('tgl_exp', 'DATE', 'ORDER BY tgl_exp ASC → logika FEFO otomatis'),
            ('stok_sisa', 'INT', 'Stok fisik tersisa per batch'),
            ('(kartu_stok)', 'JOIN', 'Jejak audit lengkap setiap mutasi stok'),
        ]
    },
]

for i, tbl in enumerate(tables):
    tx = sx + i * (half + gap)
    # Header block (teal, clean)
    add_rect(s13, tx, ct, half, Inches(0.52), TEAL_DARK)
    tb(s13, tx + Inches(0.12), ct + Inches(0.07), half - Inches(0.2), Inches(0.42),
       tbl['title'], size=11.5, bold=True, color=WHITE)
    # Description
    tb(s13, tx, ct + Inches(0.58), half, Inches(0.45),
       tbl['desc'], size=10, color=GRAY, italic=True)
    # Fields
    fy = ct + Inches(1.10)
    for field, dtype, desc in tbl['fields']:
        bg = WHITE if tbl['fields'].index((field, dtype, desc)) % 2 == 0 else RGBColor(0xF5, 0xF9, 0xFA)
        add_rect(s13, tx, fy, half, Inches(0.90), bg)
        add_rect(s13, tx, fy, Inches(0.04), Inches(0.90), TEAL)
        # Field name
        tb(s13, tx + Inches(0.14), fy + Inches(0.06), Inches(2.40), Inches(0.38),
           field, size=10.5, bold=True, color=TEAL_DARK)
        # Dtype badge
        add_rect(s13, tx + Inches(0.14), fy + Inches(0.48), Inches(0.75), Inches(0.28), TEAL_LIGHT)
        tb(s13, tx + Inches(0.14), fy + Inches(0.48), Inches(0.75), Inches(0.28),
           dtype, size=8.5, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)
        # Description
        tb(s13, tx + Inches(0.96), fy + Inches(0.50), half - Inches(1.06), Inches(0.30),
           desc, size=9.5, color=GRAY)
        fy += Inches(0.97)


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 14: KESIMPULAN
# ═══════════════════════════════════════════════════════════════════════════════
s14 = new_slide()
fill_bg(s14, CREAM)
slide_header(s14, 'Kesimpulan')

ct = content_top() + Inches(0.30)
# Large quote-style text
kesimpulan = (
    'Melalui penelitian ini, dirancang sebuah sistem informasi inventori omnichannel berbasis web yang '
    'menjembatani kesenjangan antara transaksi kasir fisik (POS) dan penjualan daring di ZenCare Medical. '
    'Dengan metode Rapid Application Development (RAD), sistem ini menjadi solusi atas permasalahan overselling '
    'dan inefisiensi konversi multi-satuan.\n\n'
    'Implementasi Gatekeeper berlogika First Expired First Out (FEFO) serta pemisahan wewenang (RBAC) yang tegas '
    'antara Admin Operasional dan Superadmin diproyeksikan menciptakan tata kelola logistik medis yang akurat, aman, dan transparan.'
)
tb(s14, Inches(0.65), ct, Inches(12.0), Inches(2.80),
   kesimpulan, size=13, color=DARK)

# 4 achievement badges
badges = [
    ('✔', 'Omnichannel Terpadu', 'POS + E-Commerce satu basis data'),
    ('✔', 'FEFO & Serial Number', 'Logistik medis tervalidasi'),
    ('✔', 'RBAC Tegas', 'Pemisahan wewenang Admin & Superadmin'),
    ('✔', 'Dashboard Manajerial', 'Pantau Harga Beli & Kartu Stok'),
]
bw = Inches(2.88)
gap = Inches(0.14)
bx = Inches(0.65)
by = ct + Inches(3.20)
for icon, title, sub in badges:
    add_rect(s14, bx, by, bw, Inches(1.40), WHITE)
    add_rect(s14, bx, by, bw, Inches(0.06), TEAL)
    tb(s14, bx + Inches(0.12), by + Inches(0.12), Inches(0.40), Inches(0.55),
       icon, size=22, bold=True, color=TEAL)
    tb(s14, bx + Inches(0.58), by + Inches(0.12), bw - Inches(0.70), Inches(0.45),
       title, size=12, bold=True, color=TITLE_CLR)
    tb(s14, bx + Inches(0.58), by + Inches(0.58), bw - Inches(0.70), Inches(0.45),
       sub, size=10, color=GRAY)
    bx += bw + gap


# ═══════════════════════════════════════════════════════════════════════════════
# SLIDE 15: PENUTUP
# ═══════════════════════════════════════════════════════════════════════════════
s15 = new_slide()
fill_bg(s15, CREAM)

# Large teal top bar
add_rect(s15, 0, 0, W, Inches(0.45), TEAL_DARK)
# Bottom teal bar
add_rect(s15, 0, Inches(7.05), W, Inches(0.45), TEAL_DARK)

# Centered content
tb(s15, Inches(0.65), Inches(1.60), Inches(12.0), Inches(1.30),
   'TERIMA KASIH', size=60, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)

add_rect(s15, Inches(4.0), Inches(3.10), Inches(5.35), Inches(0.06), TEAL)

tb(s15, Inches(0.65), Inches(3.35), Inches(12.0), Inches(0.70),
   'Rancang Bangun Sistem Informasi Inventory Omnichannel Alat Kesehatan dan Obat',
   size=15, italic=True, color=GRAY, align=PP_ALIGN.CENTER)

tb(s15, Inches(0.65), Inches(4.20), Inches(12.0), Inches(0.65),
   'Sesi Tanya Jawab (Q&A)', size=24, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)

tb(s15, Inches(0.65), Inches(5.10), Inches(12.0), Inches(0.45),
   'Marcell Chandra Kenchana  ·  322310015  ·  Universitas Ma Chung',
   size=13, color=GRAY, align=PP_ALIGN.CENTER)

# Logo area
add_rect(s15, Inches(5.55), Inches(5.75), Inches(2.25), Inches(0.65), TEAL_LIGHT)
tb(s15, Inches(5.60), Inches(5.80), Inches(1.05), Inches(0.55),
   'MA CHUNG', size=8, bold=True, color=TEAL_DARK, align=PP_ALIGN.CENTER)
tb(s15, Inches(6.70), Inches(5.80), Inches(1.00), Inches(0.55),
   'Sistem\nInformasi', size=7, color=TEAL_DARK, align=PP_ALIGN.CENTER)


# ─── SAVE ─────────────────────────────────────────────────────────────────────
prs.save(OUTPUT)
print(f'[OK] PPT Canva-style berhasil dibuat: {OUTPUT}')
print(f'[OK] Total slides: {len(prs.slides)}')
