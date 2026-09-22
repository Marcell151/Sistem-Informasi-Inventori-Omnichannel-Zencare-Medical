from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor

prs = Presentation(r'c:\xampp\htdocs\inventory_zencare\lain\File PPT\PPT TA.pptx')

# Inspect slide 1 in detail
slide = prs.slides[0]
print('=== Slide 1 Detail ===')
for shape in slide.shapes:
    print(f'Shape: {shape.name} | type={shape.shape_type}')
    print(f'  pos: left={shape.left/914400:.2f}", top={shape.top/914400:.2f}"')
    print(f'  size: width={shape.width/914400:.2f}", height={shape.height/914400:.2f}"')
    if shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                text = run.text[:60]
                font = run.font
                size = font.size.pt if font.size else 'inherit'
                bold = font.bold
                try:
                    color = font.color.rgb if font.color and font.color.type else 'inherit'
                except:
                    color = 'inherit'
                print(f'  Run: "{text}" | size={size}, bold={bold}, color={color}')

# Also check slide 2 for content layout
print()
print('=== Slide 2 Detail ===')
slide2 = prs.slides[1]
for shape in slide2.shapes:
    print(f'Shape: {shape.name} | type={shape.shape_type}')
    print(f'  pos: left={shape.left/914400:.2f}", top={shape.top/914400:.2f}"')
    print(f'  size: width={shape.width/914400:.2f}", height={shape.height/914400:.2f}"')
    if shape.has_text_frame:
        for i, para in enumerate(shape.text_frame.paragraphs[:3]):
            if para.text.strip():
                print(f'  Para {i}: "{para.text[:80]}"')
                for run in para.runs:
                    font = run.font
                    size = font.size.pt if font.size else 'inherit'
                    bold = font.bold
                    try:
                        color = font.color.rgb if font.color and font.color.type else 'inherit'
                    except:
                        color = 'inherit'
                    print(f'    Run: "{run.text[:50]}" size={size} bold={bold} color={color}')

# Check background colors/fills of shapes in slide 1
print()
print('=== Background Color check ===')
for shape in prs.slides[0].shapes:
    try:
        fill = shape.fill
        print(f'{shape.name}: fill.type={fill.type}')
        if hasattr(fill, 'fore_color'):
            try:
                print(f'  fore_color={fill.fore_color.rgb}')
            except:
                pass
    except:
        pass
