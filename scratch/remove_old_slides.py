"""
Post-process: Remove old slides from the PPT (slides 2-8 = indices 1-7)
Keep slide 1 (cover) + all newly added slides (indices 8-21)
"""
from pptx import Presentation
from pptx.oxml.ns import qn
from lxml import etree
import copy

DEST = r'c:\xampp\htdocs\inventory_zencare\lain\File PPT\Presentasi_Sempro_Marcell.pptx'

prs = Presentation(DEST)
print('Before:', len(prs.slides), 'slides')

# We need to remove slides at indices 1 through 7 (the original slides 2-8)
# python-pptx doesn't have a remove_slide method, we need to do it via XML

def delete_slide(prs, index):
    """Remove a slide from the presentation by index."""
    xml_slides = prs.slides._sldIdLst
    slides_list = list(xml_slides)
    
    # Get the rId of this slide
    slide = prs.slides[index]
    slide_part = slide.part
    
    # Find and remove from sldIdLst
    slide_elem = slides_list[index]
    xml_slides.remove(slide_elem)
    
    # Drop the relationship from the presentation
    rId = prs.slides._sldIdLst  # just for ref
    # Actually find the rId by checking part rels
    for rel in prs.part.rels.values():
        if rel.reltype == 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/slide':
            if rel._target == slide_part:
                prs.part.drop_rel(rel.rId)
                break

# Remove old slides (indices 1 to 7) - do it in reverse to not mess up indices
for i in range(7, 0, -1):
    delete_slide(prs, i)

print('After delete:', len(prs.slides), 'slides')
for i, slide in enumerate(prs.slides):
    texts = [s.text_frame.text.strip()[:50] for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    print('  Slide ' + str(i+1) + ': ' + (texts[0] if texts else '[empty]'))

prs.save(DEST)
print('Saved!')
