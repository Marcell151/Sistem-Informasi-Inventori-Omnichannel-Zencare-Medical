from pptx import Presentation
import os

prs = Presentation(r'c:\xampp\htdocs\inventory_zencare\lain\File PPT\Presentasi_Sempro_Marcell.pptx')
print('Total slides:', len(prs.slides))
fsize = os.path.getsize(r'c:\xampp\htdocs\inventory_zencare\lain\File PPT\Presentasi_Sempro_Marcell.pptx')
print('File size:', fsize, 'bytes ({:.1f} KB)'.format(fsize/1024))
for i, slide in enumerate(prs.slides):
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()[:60]
            if t:
                texts.append(t)
    first = texts[0] if texts else '[no text]'
    print('Slide ' + str(i+1) + ': ' + first)
